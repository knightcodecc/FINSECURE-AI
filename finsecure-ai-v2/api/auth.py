"""
auth.py — Authentication, Authorization, User DB, and Excel Export for FinSecure AI v2.
========================================================================================
- Separate SQLite database: data/finsecure_users.db
- Salted bcrypt password hashing (never logs or stores plaintext)
- Signed JWT tokens (HS256, 8-hour expiry)
- Rate limiting on /auth/login (5 failed attempts per 15 mins per email+IP)
- Automatic Excel export via openpyxl: data/users_export.xlsx (password_hash excluded)
  with retry & tempfile fallback for Windows file locks
"""

import os
import re
import sys
import time
import shutil
import sqlite3
import tempfile
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Dict, List, Any, Optional

import bcrypt
import jwt
from fastapi import Request, Depends, HTTPException, status, WebSocket
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Project paths
PROJECT_ROOT = Path(__file__).resolve().parent.parent


def get_data_dir() -> Path:
    """Return configured DATA_DIR (default: project_root/data) ensuring directory exists."""
    raw = os.getenv("DATA_DIR", "data")
    p = Path(raw)
    if not p.is_absolute():
        p = (PROJECT_ROOT / p).resolve()
    p.mkdir(parents=True, exist_ok=True)
    return p


def get_users_db_path() -> Path:
    """Dynamically return users SQLite database path under DATA_DIR."""
    return get_data_dir() / "finsecure_users.db"


def get_excel_export_path() -> Path:
    """Dynamically return users Excel export path under DATA_DIR."""
    return get_data_dir() / "users_export.xlsx"


USERS_DB_PATH = get_users_db_path()
EXCEL_EXPORT_PATH = get_excel_export_path()

# Environment variables
SECRET_KEY = os.environ.get("SECRET_KEY", "finsecure_jwt_secret_key_v2_2026_change_in_prod")
ALGORITHM = "HS256"
ACCESS_TOKEN_EXPIRE_HOURS = 8

ADMIN_EMAIL = os.environ.get("ADMIN_EMAIL", "admin@finsecure.ai")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "AdminSecure2026!")

# Rate limiting storage: (email, ip) -> list of failed attempt timestamps (time.time())
FAILED_ATTEMPTS: Dict[tuple, List[float]] = defaultdict(list)
RATE_LIMIT_MAX_FAILURES = 5
RATE_LIMIT_WINDOW_SECONDS = 15 * 60  # 15 minutes

EMAIL_REGEX = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
bearer_scheme = HTTPBearer(auto_error=False)


def get_users_db_connection(db_path: Optional[Path] = None) -> sqlite3.Connection:
    """Return an SQLite connection configured with WAL journal mode."""
    if db_path is None:
        db_path = get_users_db_path()
    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(db_path)
    conn.execute("PRAGMA journal_mode=WAL;")
    conn.execute("PRAGMA foreign_keys=ON;")
    return conn


def init_users_db(db_path: Optional[Path] = None):
    """Initialize users and login_events tables and bootstrap the first admin user."""
    if db_path is None:
        db_path = get_users_db_path()
    with get_users_db_connection(db_path) as conn:
        with conn:
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS users (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    full_name TEXT NOT NULL,
                    email TEXT UNIQUE NOT NULL,
                    password_hash TEXT NOT NULL,
                    role TEXT NOT NULL,
                    created_at TEXT NOT NULL,
                    last_login TEXT
                );
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS login_events (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    user_id INTEGER,
                    email TEXT NOT NULL,
                    success INTEGER NOT NULL,
                    ip TEXT,
                    timestamp TEXT NOT NULL,
                    FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE SET NULL
                );
                """
            )

        # Check if admin user exists; bootstrap if absent
        if ADMIN_EMAIL and ADMIN_PASSWORD:
            cursor = conn.cursor()
            cursor.execute("SELECT id FROM users WHERE email = ?", (ADMIN_EMAIL.strip().lower(),))
            row = cursor.fetchone()
            if not row:
                pw_hash = hash_password(ADMIN_PASSWORD)
                now_iso = datetime.now(timezone.utc).isoformat()
                with conn:
                    conn.execute(
                        """
                        INSERT INTO users (full_name, email, password_hash, role, created_at, last_login)
                        VALUES (?, ?, ?, ?, ?, ?)
                        """,
                        ("System Administrator", ADMIN_EMAIL.strip().lower(), pw_hash, "admin", now_iso, None),
                    )
                # Export initial Excel state
                try:
                    export_users_to_excel(db_path=db_path)
                except Exception as e:
                    print(f"[AUTH INIT EXCEL WARNING]: {e}")


def hash_password(plain_password: str) -> str:
    """Hash plaintext password with bcrypt salt."""
    salt = bcrypt.gensalt(rounds=12)
    return bcrypt.hashpw(plain_password.encode("utf-8"), salt).decode("utf-8")


def verify_password(plain_password: str, hashed_password: str) -> bool:
    """Verify password against bcrypt hash."""
    try:
        return bcrypt.checkpw(plain_password.encode("utf-8"), hashed_password.encode("utf-8"))
    except Exception:
        return False


def validate_password_strength(password: str):
    """Validate password length >= 8 characters."""
    if not password or len(password) < 8:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Password must be at least 8 characters long."
        )


def validate_email_format(email: str):
    """Validate standard email syntax."""
    if not email or not EMAIL_REGEX.match(email.strip()):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Invalid email address format."
        )


def check_rate_limit(email: str, ip: str):
    """Check if failed login threshold has been exceeded for email+ip in the last 15 minutes."""
    key = (email.strip().lower(), ip.strip() if ip else "unknown")
    now = time.time()
    # Filter attempts within the window
    attempts = [t for t in FAILED_ATTEMPTS[key] if now - t < RATE_LIMIT_WINDOW_SECONDS]
    FAILED_ATTEMPTS[key] = attempts
    if len(attempts) >= RATE_LIMIT_MAX_FAILURES:
        wait_mins = int((RATE_LIMIT_WINDOW_SECONDS - (now - attempts[0])) / 60) + 1
        raise HTTPException(
            status_code=status.HTTP_429_TOO_MANY_REQUESTS,
            detail=f"Too many failed login attempts. Please try again after {wait_mins} minutes."
        )


def record_failed_login(email: str, ip: str):
    """Record a failed login attempt timestamp."""
    key = (email.strip().lower(), ip.strip() if ip else "unknown")
    FAILED_ATTEMPTS[key].append(time.time())


def clear_failed_logins(email: str, ip: str):
    """Clear failed login attempts upon successful login."""
    key = (email.strip().lower(), ip.strip() if ip else "unknown")
    if key in FAILED_ATTEMPTS:
        del FAILED_ATTEMPTS[key]


def create_access_token(data: Dict[str, Any], expires_delta: Optional[timedelta] = None) -> str:
    """Generate signed JWT token with 8-hour expiry."""
    to_encode = data.copy()
    expire = datetime.now(timezone.utc) + (expires_delta or timedelta(hours=ACCESS_TOKEN_EXPIRE_HOURS))
    to_encode.update({"exp": expire})
    encoded_jwt = jwt.encode(to_encode, SECRET_KEY, algorithm=ALGORITHM)
    return encoded_jwt


def decode_access_token(token: str) -> Dict[str, Any]:
    """Decode and validate signed JWT token."""
    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
        return payload
    except jwt.ExpiredSignatureError:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Session token has expired. Please log in again."
        )
    except jwt.InvalidTokenError:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Invalid authentication token."
        )


def log_login_event(user_id: Optional[int], email: str, success: bool, ip: str, db_path: Optional[Path] = None):
    """Record authentication attempt to login_events table."""
    if db_path is None:
        db_path = get_users_db_path()
    now_iso = datetime.now(timezone.utc).isoformat()
    try:
        with get_users_db_connection(db_path) as conn:
            with conn:
                conn.execute(
                    """
                    INSERT INTO login_events (user_id, email, success, ip, timestamp)
                    VALUES (?, ?, ?, ?, ?)
                    """,
                    (user_id, email, 1 if success else 0, ip, now_iso),
                )
    except Exception as e:
        print(f"[AUTH LOG EVENT ERROR]: {e}")


async def get_current_user(
    request: Request,
    credentials: Optional[HTTPAuthorizationCredentials] = Depends(bearer_scheme)
) -> Dict[str, Any]:
    """FastAPI dependency to extract and verify the current authenticated user from Authorization header."""
    token = None
    if credentials:
        token = credentials.credentials
    else:
        # Check query param for fallback
        token = request.query_params.get("token")

    if not token:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Authentication required. Please provide a valid Bearer token."
        )

    payload = decode_access_token(token)
    user_email = payload.get("sub")
    if not user_email:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Invalid token payload."
        )

    # Verify user still exists in database
    with get_users_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT id, full_name, email, role FROM users WHERE email = ?", (user_email.lower(),))
        row = cursor.fetchone()
        if not row:
            raise HTTPException(
                status_code=status.HTTP_401_UNAUTHORIZED,
                detail="User account no longer exists."
            )
        return {
            "id": row[0],
            "full_name": row[1],
            "email": row[2],
            "role": row[3],
        }


async def get_current_admin(current_user: Dict[str, Any] = Depends(get_current_user)) -> Dict[str, Any]:
    """Dependency verifying the caller has admin privileges."""
    if current_user.get("role") != "admin":
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Admin privileges required for this action."
        )
    return current_user


def export_users_to_excel(db_path: Optional[Path] = None, export_path: Optional[Path] = None):
    """
    Export SQLite users and login_events to Excel (.xlsx) using openpyxl.
    - Sheet 1: Users (password_hash EXCLUDED)
    - Sheet 2: Login Events
    - Windows-safe: Handles PermissionError gracefully if file is open in Excel.
    """
    if db_path is None:
        db_path = get_users_db_path()
    if export_path is None:
        export_path = get_excel_export_path()
    export_path.parent.mkdir(parents=True, exist_ok=True)

    with get_users_db_connection(db_path) as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT id, full_name, email, role, created_at, last_login FROM users ORDER BY id ASC")
        users_rows = cursor.fetchall()

        cursor.execute("SELECT id, user_id, email, success, ip, timestamp FROM login_events ORDER BY id DESC")
        login_rows = cursor.fetchall()

    wb = openpyxl.Workbook()

    # Sheet 1: Users
    ws_users = wb.active
    ws_users.title = "Users"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    user_headers = ["ID", "Full Name", "Email Address", "Role", "Created At (UTC)", "Last Login (UTC)"]
    ws_users.append(user_headers)
    for col_num in range(1, len(user_headers) + 1):
        cell = ws_users.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for r in users_rows:
        ws_users.append(list(r))

    for row in ws_users.iter_rows(min_row=2, max_row=len(users_rows) + 1, min_col=1, max_col=len(user_headers)):
        for cell in row:
            cell.border = thin_border
            cell.alignment = left_align

    # Sheet 2: Login Events
    ws_events = wb.create_sheet(title="Login Events")
    event_headers = ["Event ID", "User ID", "Email", "Status", "IP Address", "Timestamp (UTC)"]
    ws_events.append(event_headers)
    for col_num in range(1, len(event_headers) + 1):
        cell = ws_events.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
        cell.alignment = center_align

    for r in login_rows:
        status_text = "SUCCESS" if r[3] == 1 else "FAILED"
        ws_events.append([r[0], r[1], r[2], status_text, r[4] or "unknown", r[5]])

    for row in ws_events.iter_rows(min_row=2, max_row=len(login_rows) + 1, min_col=1, max_col=len(event_headers)):
        for cell in row:
            cell.border = thin_border
            cell.alignment = left_align

    # Auto-adjust column widths
    for ws in [ws_users, ws_events]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Save safely with retry and tempfile fallback for Windows PermissionError
    temp_file = export_path.parent / f"users_export_{int(time.time()*1000)}.tmp.xlsx"
    try:
        wb.save(temp_file)
        # Attempt atomic replace
        max_retries = 3
        for attempt in range(max_retries):
            try:
                if export_path.exists():
                    os.replace(temp_file, export_path)
                else:
                    shutil.move(temp_file, export_path)
                break
            except PermissionError:
                if attempt == max_retries - 1:
                    print(f"[EXCEL EXPORT]: File {export_path} is currently locked by Excel. Saved update to {temp_file}")
                    return
                time.sleep(0.1)
    except Exception as e:
        print(f"[EXCEL EXPORT EXCEPTION]: {e}")
        if temp_file.exists():
            try:
                temp_file.unlink()
            except Exception:
                pass
